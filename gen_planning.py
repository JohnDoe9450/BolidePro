#!/usr/bin/env python3
# -*- coding: utf-8 -*-
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BLUE = RGBColor(0x1F, 0x4E, 0x9C)
DARK = RGBColor(0x22, 0x22, 0x22)
GREY = RGBColor(0x66, 0x66, 0x66)
HEADER_BG = "1F4E9C"
ALT_BG = "EEF2FA"
RED_BG = "FBEAEA"

# Lignes mises en évidence en rouge dans le tableau (par numéro de semaine)
RED_WEEKS = {"S1", "S2", "S8"}


def new_doc():
    doc = Document()
    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(10.5)
    style.font.color.rgb = DARK
    return doc


def shade_cell(cell, color_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color_hex)
    tcPr.append(shd)


def set_cell_text(cell, text, bold=False, color=None, size=10, align=None):
    cell.text = ''
    p = cell.paragraphs[0]
    if align:
        p.alignment = align
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color
    return cell


def add_heading(doc, text, size=16, color=BLUE, space_before=14, space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(size)
    run.font.color.rgb = color
    return p


def add_sub(doc, text, size=12.5, color=BLUE):
    return add_heading(doc, text, size=size, color=color, space_before=12, space_after=4)


def add_title_block(doc, subtitle):
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = title.add_run('Bolide Pro')
    r.bold = True
    r.font.size = Pt(30)
    r.font.color.rgb = BLUE

    sub = doc.add_paragraph()
    r = sub.add_run(subtitle)
    r.font.size = Pt(15)
    r.font.color.rgb = GREY

    meta = doc.add_paragraph()
    r = meta.add_run('Démarrage : lundi 15 juin 2026 — Gabon / Afrique centrale (XAF)')
    r.italic = True
    r.font.size = Pt(10)
    r.font.color.rgb = GREY
    doc.add_paragraph()


# ---------- DONNÉES ----------
ROWS = [
    ("Semaine", "Dates", "Phase", "Livrable"),
    ("S1", "15 – 19 juin 2026", "Fleet OS V1.5",
     "Gestion des véhicules opérationnelle : chaque flotte peut enregistrer ses véhicules, "
     "suivre leur état (actif / immobilisé) et associer un chauffeur à un véhicule. "
     "Un chauffeur ne peut conduire qu'un véhicule à la fois, et un véhicule n'a qu'un seul "
     "chauffeur actif."),
    ("S2", "22 – 26 juin 2026", "Fleet OS V1.5",
     "Suivi des incidents et des dépenses : déclaration et résolution d'un incident de bout en "
     "bout (de la déclaration à la clôture), enregistrement des dépenses (carburant, réparations), "
     "détection automatique des retards et blocages de chauffeurs chaque soir, et export des "
     "données en Excel/CSV."),
    ("S3", "29 juin – 3 juil. 2026", "Bolide Pro",
     "Affectations dédiées : un manager d'entreprise se voit attribuer des chauffeurs en "
     "exclusivité, avec leur taux journalier et leurs horaires contractuels. Le système garantit "
     "qu'un chauffeur n'a qu'une seule affectation active et conserve l'historique de toutes les "
     "modifications."),
    ("S4", "6 – 10 juil. 2026", "Bolide Pro",
     "Validation journalière des activités : le chauffeur déclare sa fin de journée, le manager "
     "valide ou refuse. Seules les journées validées seront comptabilisées. Chaque partie est "
     "notifiée en temps réel (validation, refus, commentaire)."),
    ("S5", "13 – 17 juil. 2026", "Bolide Pro",
     "Gestion des heures supplémentaires et calcul du salaire : les heures sup sont détectées "
     "automatiquement dès qu'un chauffeur dépasse son horaire, puis approuvées (totalement ou "
     "partiellement) par le manager. Le salaire mensuel est calculé et mis à jour en temps réel "
     "à chaque validation."),
    ("S6", "20 – 24 juil. 2026", "Bolide Pro",
     "Tableau de bord manager : une vue unique et consolidée de tous ses chauffeurs — activité, "
     "validations en attente, heures travaillées, coût mensuel et performance. Conçu pour une "
     "consultation rapide même avec un grand nombre de chauffeurs."),
    ("S7", "27 – 31 juil. 2026", "Bolide Pro",
     "Espace chauffeur et notation : chaque chauffeur dispose de son espace personnel (salaire "
     "estimé, journées, notifications, déclaration de problème). Le manager peut noter ses "
     "chauffeurs (1 à 5), avec un score moyen affiché. Notifications par email et in-app finalisées."),
    ("S8", "3 – 7 août 2026", "Tests & corrections",
     "Recette finale : tests complets de tous les parcours (validation, salaire, sécurité, "
     "cloisonnement des données entre clients), correction des anomalies détectées et "
     "préparation de la mise en production. Aucune nouvelle fonctionnalité — uniquement "
     "fiabilisation."),
]

TECH = [
    ("S1 — 15 au 19 juin 2026 — Fleet OS V1.5 : Véhicules & Affectations [4.7]", [
        "Entité `vehicules` : CRUD, champs (flotte_id, type, immatriculation, statut).",
        "Statuts véhicule : actif | immobilisé (hors_service = V2). Passage en immobilisé → "
        "fin automatique de l'affectation + alerte dashboard.",
        "Entité `affectations_vehicule` (chauffeur ↔ véhicule), historisée.",
        "Contrainte : 1 affectation active max par chauffeur ET 1 par véhicule.",
        "Index DB sur fleet_id, driver_id, status.",
    ]),
    ("S2 — 22 au 26 juin 2026 — Fleet OS V1.5 : Incidents, Dépenses, Automatismes [4.8 / 9.3]", [
        "Workflow incident `incidents` : déclaré → en_analyse → en_cours → résolu → clôturé "
        "(non supprimable). Incident critique → immobilisation auto du véhicule.",
        "Entité `depenses` (carburant, réparation) liée à flotte/véhicule/incident, écriture "
        "ledger DEP au wallet.",
        "Job CRON clôture journalière 23h59 : calcul retard = attendu_cumulé − payé_cumulé, "
        "mise à jour statuts chauffeur (actif/retard/bloqué) selon seuil_alerte / seuil_blocage, "
        "snapshot wallet. Idempotent.",
        "Exports CSV/Excel (paiements, wallet) [P1].",
    ]),
    ("S3 — 29 juin au 3 juillet 2026 — Bolide Pro : Affectations dédiées [5.1]", [
        "Migrations PostgreSQL : `pro_assignments`, `pro_assignment_history`, tables Pro restantes.",
        "Nouveau rôle `pro_manager` + mise à jour matrice RBAC, middleware autorisation Pro "
        "(rôle + appartenance flotte, 403 cross-flotte).",
        "Endpoints : POST /pro/assignments, PATCH /pro/assignments/:id/status, "
        "GET /pro/assignments/:id, GET /pro/drivers/:id/assignment, "
        "GET /pro/managers/:id/assignments, GET /pro/assignments/:id/history.",
        "Contrainte métier UNIQUE(driver_id) WHERE status='active' → HTTP 409.",
        "Pas de suppression (status='suspended'). Toute modif → historique (changed_by, timestamp).",
        "vehicle_id NULLABLE : véhicule dédié optionnel (pas de dépendance dure à V1.5).",
    ]),
    ("S4 — 6 au 10 juillet 2026 — Bolide Pro : Validation journalière [5.2]", [
        "Entités `daily_logs` + `daily_validations`. UNIQUE(driver_id, date).",
        "PATCH /pro/daily-logs/:id/end-time (driver) → daily_log status='pending' + push manager.",
        "POST /pro/daily-logs/:id/validate et /reject (pro_manager, super_admin) ; "
        "/admin-override (super_admin, trace obligatoire).",
        "GET /pro/daily-logs?driver_id=&month= et GET /pro/assignments/:id/pending.",
        "Règles : seules les journées 'validated' comptent ; on ne valide/refuse qu'un 'pending' ; "
        "immuable post-validation sauf override tracé ; manager limité à SES chauffeurs.",
        "Notifications VAL (→ manager) et VAL_RES (→ driver).",
    ]),
    ("S5 — 13 au 17 juillet 2026 — Bolide Pro : Heures supplémentaires & Salaire [5.3 / 5.4]", [
        "Entité `overtime_requests`. Déclenchement auto si end_time_declared > expected_end_time : "
        "création requête (requested_minutes = DATEDIFF, status='pending') + notif manager.",
        "POST /pro/overtime/:id/approve (total ou {validated_minutes: X} partiel) ; "
        "/reject avec corrected_end_time OBLIGATOIRE → HTTP 422 sinon (validé côté serveur avant écriture).",
        "Entité `salaries`. Formule : total = (jours_validés × daily_rate) + "
        "(overtime_minutes_total × overtime_rate_per_min) + SUM(adjustments).",
        "Recalcul auto idempotent à chaque validation/refus daily_log ou overtime, et à chaque "
        "ajustement. Job asynchrone si > 200ms.",
        "Endpoints : GET /pro/drivers/:id/salary, /salary/history, POST /salary/adjustment "
        "(super_admin, montant + motif obligatoire).",
        "Notifications OT (→ manager) et OT_RES (→ driver).",
    ]),
    ("S6 — 20 au 24 juillet 2026 — Bolide Pro : Dashboard manager Pro [5.5]", [
        "Endpoint agrégé unique GET /pro/managers/:id/dashboard?month=YYYY-MM (objectif < 800ms "
        "pour 50 chauffeurs).",
        "Sections : Opérationnel (chauffeurs actifs, validations en attente, incidents ouverts), "
        "Performance (note moyenne 30j, top chauffeurs), Temps de travail (heures travaillées, "
        "heures sup en attente), Financier (coût total, coût heures sup).",
        "Cache lecture Redis (TTL court 30s–2min). Index DB adaptés.",
    ]),
    ("S7 — 27 au 31 juillet 2026 — Bolide Pro : Interface Driver & Notation [5.6 / 5.7]", [
        "Écrans driver Pro : dashboard personnel (salaire estimé, journées validées/en attente, "
        "heures sup), Mes journées (déclarer fin de journée), Mon salaire (lecture seule), "
        "Notifications, Déclarer un problème (→ crée Incident), Support.",
        "Entité `ratings` : POST /pro/ratings, GET /pro/drivers/:id/ratings. Score 1–5, "
        "CHECK(1..5), commentaire recommandé si < 3.",
        "Immuabilité : aucun PUT/PATCH/DELETE sur ratings. Score global = AVG(score) sur 30j.",
        "Manager limité à SES affectations actives (vérif server-side).",
        "Finalisation notifications Pro (in-app + email) ; règle UX : notification actionnable "
        "(lien direct vers l'entité).",
    ]),
    ("S8 — 3 au 7 août 2026 — Tests & corrections [12]", [
        "Tests unitaires : calcul retard, calcul salaire, formule overtime, validation statuts.",
        "Tests d'intégration : workflow paiement → wallet → ledger ; workflow daily_log → salary.",
        "Tests de sécurité : cross-fleet access (403 attendu), rôle insuffisant, suppression "
        "interdite (ledger, ratings, daily_validations).",
        "Vérification des critères d'acceptation : HTTP 409 (2e affectation active), HTTP 422 "
        "(refus OT sans corrected_end_time), recalcul salaire < 200ms / async, dashboard < 800ms.",
        "Tests de régression : l'ajout de Bolide Pro ne casse pas les calculs wallet Fleet OS.",
        "Correction des anomalies, documentation OpenAPI/Swagger à jour, préparation mise en prod.",
    ]),
]


# ---------- BLOCS RÉUTILISABLES ----------
def build_planning_table(doc, heading_text):
    add_heading(doc, heading_text, size=16)
    intro = doc.add_paragraph()
    intro.add_run(
        "Ce planning couvre la finalisation du socle Fleet OS V1.5 (semaines 1 à 2) puis le "
        "développement complet de la couche Bolide Pro (semaines 3 à 7). La huitième et dernière "
        "semaine est entièrement consacrée aux tests et corrections avant mise en production."
    )

    table = doc.add_table(rows=0, cols=4)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths = [Inches(0.6), Inches(1.5), Inches(1.1), Inches(3.8)]

    for i, (c0, c1, c2, c3) in enumerate(ROWS):
        row = table.add_row()
        cells = row.cells
        for j, txt in enumerate((c0, c1, c2, c3)):
            cells[j].width = widths[j]
            if i == 0:
                set_cell_text(cells[j], txt, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF), size=10.5)
                shade_cell(cells[j], HEADER_BG)
            else:
                set_cell_text(cells[j], txt, bold=(j <= 2),
                              color=BLUE if j <= 2 else DARK, size=10)
                if c0 in RED_WEEKS:
                    shade_cell(cells[j], RED_BG)
                elif i % 2 == 0:
                    shade_cell(cells[j], ALT_BG)

    doc.add_paragraph()
    note = doc.add_paragraph()
    r = note.add_run("Dépendances : ")
    r.bold = True
    r.font.color.rgb = BLUE
    note.add_run(
        "Fleet OS V1.5 (S1–S2) fournit le socle (véhicules, incidents) réutilisé par Bolide Pro. "
        "Au sein de Bolide Pro, chaque semaine s'appuie sur la précédente : les affectations (S3) "
        "conditionnent la validation (S4), qui conditionne le salaire (S5), lui-même agrégé dans "
        "les tableaux de bord (S6). Mise en production prévue : lundi 10 août 2026."
    ).font.size = Pt(10)


def build_tech_section(doc, heading_text):
    add_heading(doc, heading_text, size=16)
    tech_intro = doc.add_paragraph()
    tech_intro.add_run(
        "Cette section détaille, pour chaque semaine, les modules et tâches techniques à réaliser. "
        "Elle s'adresse au développeur. Références entre crochets = sections du cahier des charges v2.0."
    ).italic = True

    for header, items in TECH:
        add_sub(doc, header, size=12)
        for it in items:
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(it)
            run.font.size = Pt(10)


def add_footer(doc, label):
    doc.add_paragraph()
    foot = doc.add_paragraph()
    r = foot.add_run(label)
    r.italic = True
    r.font.size = Pt(9)
    r.font.color.rgb = GREY
    foot.alignment = WD_ALIGN_PARAGRAPH.CENTER


# ---------- FICHIER COMBINÉ ----------
doc = new_doc()
add_title_block(doc, 'Planning de développement — 8 semaines')
build_planning_table(doc, "1. Vue d'ensemble du planning")
doc.add_page_break()
build_tech_section(doc, "2. Détail technique par livrable (Modules & tâches)")
add_footer(doc, 'Bolide — Planning de développement Bolide Pro (8 semaines) — usage interne')
doc.save('/home/user/BolidePro/Bolide_Pro_Planning_8_semaines.docx')

# ---------- FICHIER 1 : PLANNING (point 1) ----------
doc1 = new_doc()
add_title_block(doc1, 'Planning de développement — Vue d\'ensemble')
build_planning_table(doc1, "Vue d'ensemble du planning")
add_footer(doc1, 'Bolide — Planning de développement Bolide Pro (8 semaines) — usage interne')
doc1.save('/home/user/BolidePro/Bolide_Pro_Planning_Vue_ensemble.docx')

# ---------- FICHIER 2 : DÉTAIL TECHNIQUE (point 2) ----------
doc2 = new_doc()
add_title_block(doc2, 'Détail technique par livrable — Modules & tâches')
build_tech_section(doc2, "Détail technique par livrable (Modules & tâches)")
add_footer(doc2, 'Bolide — Détail technique Bolide Pro (8 semaines) — usage interne')
doc2.save('/home/user/BolidePro/Bolide_Pro_Detail_technique.docx')

print('Saved 3 files.')
